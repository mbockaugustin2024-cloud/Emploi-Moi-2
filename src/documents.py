import html
import json
import os
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from supabase import create_client
from ai_writer import generate_ai_content

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
with open(
    os.path.join(BASE_DIR, "config", "profile.json"),
    "r",
    encoding="utf-8",
) as f:
    PROFILE = json.load(f)

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_SERVICE_ROLE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
supabase = create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)

BUCKET = "application-documents"
URL_TTL = 60 * 60 * 24 * 90

FRENCH_MARKERS = (
    "french", "français", "francais", "francophone",
    "support client", "téléconseiller", "conseiller client",
)


def plain(value):
    text = html.unescape(str(value or ""))
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def detect_language(job):
    raw = " ".join(
        str(job.get(k) or "")
        for k in ("title", "description", "detected_language")
    ).lower()
    return "fr" if any(x in raw for x in FRENCH_MARKERS) else "en"


def matched_skills(job):
    text = (
        plain(job.get("title"))
        + " "
        + plain(job.get("description"))
    ).lower()
    skills = []
    for skill in PROFILE["skills"]:
        tokens = [
            token
            for token in re.findall(r"[a-zà-ÿ0-9]+", skill.lower())
            if len(token) > 2
        ]
        if tokens and any(token in text for token in tokens):
            skills.append(skill)
    return skills[:8]


def relevant_experience(job):
    text = (
        plain(job.get("title"))
        + " "
        + plain(job.get("description"))
    ).lower()
    hospitality = any(
        key in text
        for key in (
            "restaurant", "waiter", "server", "hospitality",
            "hotel", "réception", "reception",
        )
    )
    experiences = PROFILE["experience"]
    if hospitality:
        return [
            e for e in experiences
            if "Waiter" in e["title"] or "Receptionist" in e["title"]
        ]
    return [
        e for e in experiences
        if any(
            key in " ".join(e["skills"]).lower()
            for key in (
                "customer", "call", "crm", "sales",
                "appointment", "supervision",
            )
        )
    ][:5]


def make_cv(job, ai_data=None):
    language = detect_language(job)
    name = PROFILE["candidate"]["name"]
    contact = PROFILE["candidate"]["contact"]
    target_role = plain(job.get("title")) or "Customer Support / Call Center"
    skills = matched_skills(job)
    experiences = relevant_experience(job)

    if language == "fr":
        summary = (
            "Professionnel francophone de la relation client et des centres "
            "d'appels, avec expérience en vente B2B, appels entrants et sortants, "
            "suivi CRM, prospection, traitement des objections et supervision."
        )
        sections = {
            "summary": "PROFIL",
            "skills": "COMPÉTENCES CLÉS",
            "experience": "EXPÉRIENCE PROFESSIONNELLE",
            "education": "FORMATION",
            "languages": "LANGUES",
        }
        language_line = (
            "Français - langue maternelle / courant | "
            "Anglais - débutant | Espagnol - débutant"
        )
        if not ai_data:
            ai_data = {}
        summary = ai_data.get("summary_fr") or summary
        ai_skills = ai_data.get("selected_skills") or []
        if ai_skills:
            skills = ai_skills[:8]
    else:
        summary = PROFILE["professional_profile"]
        sections = {
            "summary": "PROFESSIONAL PROFILE",
            "skills": "CORE SKILLS",
            "experience": "PROFESSIONAL EXPERIENCE",
            "education": "EDUCATION",
            "languages": "LANGUAGES",
        }
        language_line = (
            "French - Native / Fluent | "
            "English - Beginner | Spanish - Beginner"
        )
        if not ai_data:
            ai_data = {}
        summary = ai_data.get("summary_en") or summary
        ai_skills = ai_data.get("selected_skills") or []
        if ai_skills:
            skills = ai_skills[:8]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8.2)

    p = doc.add_paragraph()
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(16)

    p = doc.add_paragraph()
    run = p.add_run(target_role)
    run.bold = True
    run.font.size = Pt(10.5)

    doc.add_paragraph(
        f"{PROFILE['candidate']['current_location']} | "
        f"{contact['phones'][0]} / {contact['phones'][1]} | "
        f"{contact['email']}"
    )

    p = doc.add_paragraph(sections["summary"])
    p.runs[0].bold = True
    doc.add_paragraph(summary)

    p = doc.add_paragraph(sections["skills"])
    p.runs[0].bold = True
    doc.add_paragraph(" • ".join(skills or PROFILE["skills"][:8]))

    p = doc.add_paragraph(sections["experience"])
    p.runs[0].bold = True

    all_experiences = [
        {
            "title": PROFILE["current_experience"]["title"],
            "company": "",
            "location": "Tunis, Tunisia",
            "start": "2026-05",
            "end": "Present",
            "skills": PROFILE["current_experience"]["responsibilities"],
        }
    ] + experiences

    for experience in all_experiences[:6]:
        p = doc.add_paragraph()
        run = p.add_run(
            experience["title"]
            + (
                f" - {experience['company']}"
                if experience.get("company")
                else ""
            )
        )
        run.bold = True
        doc.add_paragraph(
            f"{experience.get('location', '')} | "
            f"{experience.get('start', '')} - {experience.get('end', '')}"
        )
        for item in experience["skills"][:3]:
            paragraph = doc.add_paragraph(item)
            paragraph.paragraph_format.left_indent = Inches(0.14)
            paragraph.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph(sections["education"])
    p.runs[0].bold = True
    education = PROFILE["education"][0]
    doc.add_paragraph(
        f"{education['qualification']} - {education['country']}"
    )

    p = doc.add_paragraph(sections["languages"])
    p.runs[0].bold = True
    doc.add_paragraph(language_line)

    return doc


def make_letter(job, ai_data=None):
    language = detect_language(job)
    company = plain(job.get("company")) or "the hiring team"
    title = plain(job.get("title")) or "the position"
    skills = matched_skills(job)

    if language == "fr":
        if ai_data and ai_data.get("letter_fr"):
            return [
                f"Objet : Candidature - {title}",
                ai_data["letter_fr"],
            ]
        body = [
            f"Objet : Candidature - {title}",
            (
                f"Madame, Monsieur,\n\n"
                f"Je vous adresse ma candidature pour le poste de {title} "
                f"au sein de {company}."
            ),
            (
                "Mon expérience en relation client et centres d'appels, "
                "notamment en vente B2B dans le secteur de l'énergie, m'a "
                "permis de développer une pratique solide des appels "
                "entrants et sortants, du suivi CRM, de la prospection, "
                "du traitement des objections et de la fidélisation client."
            ),
            (
                "Compétences directement pertinentes : "
                + (
                    ", ".join(skills[:6])
                    if skills
                    else "relation client, appels entrants et sortants, "
                         "suivi CRM et prospection"
                )
                + "."
            ),
            (
                "Sérieux, à l'écoute et habitué aux objectifs commerciaux, "
                "je suis disponible pour une opportunité internationale "
                "et une relocalisation lorsque celle-ci est proposée "
                "par l'employeur."
            ),
            (
                "Je serais heureux d'échanger avec vous afin de vous "
                "présenter plus précisément mon parcours et ma motivation.\n\n"
                "Cordialement,\nRodrigue Augustin Mbock"
            ),
        ]
    else:
        if ai_data and ai_data.get("letter_en"):
            return [
                f"Subject: Application - {title}",
                ai_data["letter_en"],
            ]
        body = [
            f"Subject: Application - {title}",
            (
                "Dear Hiring Team,\n\n"
                f"I am applying for the {title} position at {company}."
            ),
            (
                "I have professional experience in customer service and "
                "call centers, including B2B energy sales, inbound and "
                "outbound calls, CRM follow-up, prospecting, objection "
                "handling and customer retention."
            ),
            (
                "Relevant strengths for this position include: "
                + (
                    ", ".join(skills[:6])
                    if skills
                    else "customer service, inbound/outbound calls, "
                         "CRM follow-up and prospecting"
                )
                + "."
            ),
            (
                "I am comfortable working with clear targets, "
                "customer-facing responsibilities and follow-up processes. "
                "I am also open to international relocation when offered "
                "by the employer."
            ),
            (
                "Thank you for considering my application. I would welcome "
                "the opportunity to discuss how my experience could support "
                "your team.\n\nKind regards,\nRodrigue Augustin Mbock"
            ),
        ]
    return body


def extract_url(result):
    if isinstance(result, dict):
        return (
            result.get("signedURL")
            or result.get("signedUrl")
            or result.get("signed_url")
            or result.get("url")
        )
    return (
        getattr(result, "signedURL", None)
        or getattr(result, "signed_url", None)
    )


def ensure_bucket():
    try:
        supabase.storage.get_bucket(BUCKET)
        return
    except Exception:
        pass

    try:
        supabase.storage.create_bucket(
            BUCKET,
            options={
                "public": False,
                "allowed_mime_types": [
                    "application/pdf",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ],
            },
        )
    except Exception as exc:
        message = str(exc).lower()
        if "already exists" not in message and "duplicate" not in message:
            raise


def upload_bytes(path, data, content_type):
    supabase.storage.from_(BUCKET).upload(
        path,
        data,
        {"content-type": content_type, "upsert": "true"},
    )
    result = supabase.storage.from_(BUCKET).create_signed_url(
        path,
        URL_TTL,
    )
    url = extract_url(result)
    if not url:
        raise RuntimeError(f"Impossible de récupérer l'URL signée pour {path}")
    return url


def build_docx_bytes(doc):
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf_from_story(story):
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    document.build(story)
    return buffer.getvalue()


def build_letter_pdf(job, ai_data=None):
    body = make_letter(job, ai_data)
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SmallLetter",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            spaceAfter=7,
        )
    )
    story = []
    for index, block in enumerate(body):
        story.append(
            Paragraph(
                html.escape(block).replace("\n", "<br/>"),
                styles["Title"] if index == 0 else styles["SmallLetter"],
            )
        )
        if index == 0:
            story.append(Spacer(1, 7))
    return build_pdf_from_story(story)


def build_cv_pdf(job, ai_data=None):
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CVSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.15,
            leading=8.6,
            spaceAfter=1.5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CVHeading",
            parent=styles["Heading4"],
            fontName="Helvetica-Bold",
            fontSize=8.2,
            leading=9.2,
            spaceBefore=4,
            spaceAfter=2,
        )
    )

    language = detect_language(job)
    skills = (
        ai_data.get("selected_skills", [])
        if ai_data
        else []
    ) or matched_skills(job) or PROFILE["skills"][:8]

    if language == "fr":
        summary = (
            (ai_data or {}).get("summary_fr")
            or "Professionnel francophone de la relation client et des centres "
               "d'appels, avec expérience en vente B2B, appels entrants et "
               "sortants, suivi CRM, prospection, traitement des objections "
               "et supervision."
        )
        language_line = (
            "Français - langue maternelle / courant | "
            "Anglais - débutant | Espagnol - débutant"
        )
        skills_heading = "COMPÉTENCES CLÉS"
        experience_heading = "EXPÉRIENCE PROFESSIONNELLE"
        education_heading = "FORMATION"
        languages_heading = "LANGUES"
    else:
        summary = (
            (ai_data or {}).get("summary_en")
            or PROFILE["professional_profile"]
        )
        language_line = (
            "French - Native / Fluent | "
            "English - Beginner | Spanish - Beginner"
        )
        skills_heading = "CORE SKILLS"
        experience_heading = "PROFESSIONAL EXPERIENCE"
        education_heading = "EDUCATION"
        languages_heading = "LANGUAGES"

    story = [
        Paragraph(
            html.escape(PROFILE["candidate"]["name"]),
            styles["Title"],
        ),
        Paragraph(
            html.escape(
                plain(job.get("title"))
                or "Customer Support / Call Center"
            ),
            styles["Heading3"],
        ),
        Paragraph(
            html.escape(
                f"{PROFILE['candidate']['current_location']} | "
                f"{PROFILE['candidate']['contact']['email']} | "
                f"{PROFILE['candidate']['contact']['phones'][0]} / "
                f"{PROFILE['candidate']['contact']['phones'][1]}"
            ),
            styles["CVSmall"],
        ),
        Paragraph(html.escape(summary), styles["CVSmall"]),
        Paragraph(html.escape(skills_heading), styles["CVHeading"]),
        Paragraph(
            html.escape(" • ".join(skills[:8])),
            styles["CVSmall"],
        ),
        Paragraph(html.escape(experience_heading), styles["CVHeading"]),
    ]

    experiences = [
        {
            "title": PROFILE["current_experience"]["title"],
            "company": "",
            "location": "Tunis, Tunisia",
            "start": "2026-05",
            "end": "Present",
            "skills": PROFILE["current_experience"]["responsibilities"],
        }
    ] + relevant_experience(job)

    for experience in experiences[:6]:
        story.append(
            Paragraph(
                html.escape(
                    experience["title"]
                    + (
                        " - " + experience["company"]
                        if experience.get("company")
                        else ""
                    )
                ),
                styles["Heading4"],
            )
        )
        story.append(
            Paragraph(
                html.escape(
                    " | ".join(
                        item
                        for item in (
                            experience.get("location"),
                            experience.get("start"),
                            experience.get("end"),
                        )
                        if item
                    )
                ),
                styles["CVSmall"],
            )
        )
        for item in experience["skills"][:3]:
            story.append(
                Paragraph("- " + html.escape(item), styles["CVSmall"])
            )

    education = PROFILE["education"][0]
    story.extend(
        [
            Paragraph(html.escape(education_heading), styles["CVHeading"]),
            Paragraph(
                html.escape(
                    f"{education['qualification']} - {education['country']}"
                ),
                styles["CVSmall"],
            ),
            Paragraph(html.escape(languages_heading), styles["CVHeading"]),
            Paragraph(html.escape(language_line), styles["CVSmall"]),
        ]
    )

    return build_pdf_from_story(story)


def generate_for_job(job, ai_data=None):
    ensure_bucket()

    job_key = re.sub(
        r"[^a-z0-9]+",
        "-",
        str(job.get("id") or job.get("source_job_id") or "job"),
        flags=re.I,
    ).strip("-") or "job"

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    folder = f"jobs/{job_key}-{stamp}"

    cv_path = f"{folder}/CV_Rodrigue_Mbock.pdf"
    cv_docx_path = f"{folder}/CV_Rodrigue_Mbock.docx"
    letter_path = f"{folder}/Lettre_Rodrigue_Mbock.pdf"

    cv_doc = make_cv(job, ai_data)

    cv_url = upload_bytes(
        cv_path,
        build_cv_pdf(job, ai_data),
        "application/pdf",
    )

    upload_bytes(
        cv_docx_path,
        build_docx_bytes(cv_doc),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    letter_url = upload_bytes(
        letter_path,
        build_letter_pdf(job, ai_data),
        "application/pdf",
    )

    return cv_url, letter_url


def main():
    cutoff = (datetime.now(timezone.utc) - timedelta(days=90)).isoformat()
    rows = (
        supabase
        .table("jobs")
        .select("*")
        .eq("is_active", True)
        .gte("score", 40)
        .gte("publication_date", cutoff)
        .limit(40)
        .execute()
        .data
        or []
    )

    rows = [
        row
        for row in rows
        if not row.get("cv_url") or not row.get("cover_letter_url")
    ]

    print(f"Documents à générer: {len(rows)}")

    for row in rows:
        try:
            ai_data = generate_ai_content(row, PROFILE)
            if ai_data:
                print(f"IA active pour: {row.get('title')}")
            else:
                print(f"Fallback automatique pour: {row.get('title')}")
            cv_url, letter_url = generate_for_job(row, ai_data)
            (
                supabase
                .table("jobs")
                .update({
                    "cv_url": cv_url,
                    "cover_letter_url": letter_url,
                    "cv_text": (
                        "CV généré à partir du profil maître ; adaptation IA si la clé OpenRouter est configurée."
                        "au contenu de l'offre."
                    ),
                    "cover_letter_text": (
                        "Lettre générée à partir du profil maître ; adaptation IA si la clé OpenRouter est configurée."
                        "adaptée au contenu de l'offre."
                    ),
                })
                .eq("id", row["id"])
                .execute()
            )
            print(f"Documents générés: {row.get('title')}")
        except Exception as exc:
            print(f"Erreur documents {row.get('title')}: {exc}")


if __name__ == "__main__":
    main()
